from __future__ import annotations

from pathlib import Path
from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


ROOT = Path(__file__).resolve().parents[1]
OUTPUT_DIR = ROOT / "deliverables"
OUTPUT_PATH = OUTPUT_DIR / "video-agent-studio-workflow-explainer.docx"


def set_cell_shading(cell, fill: str) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:fill"), fill)
    tc_pr.append(shd)


def set_document_defaults(document: Document) -> None:
    styles = document.styles
    normal = styles["Normal"]
    normal.font.name = "Arial"
    normal.font.size = Pt(10.5)
    title = styles["Title"]
    title.font.name = "Arial"
    title.font.size = Pt(22)
    title.font.bold = True
    for style_name in ("Heading 1", "Heading 2", "Heading 3"):
        style = styles[style_name]
        style.font.name = "Arial"
        style.font.bold = True
    if "No Spacing" in styles:
        styles["No Spacing"].font.name = "Courier New"
        styles["No Spacing"].font.size = Pt(9)


def add_title(document: Document) -> None:
    title = document.add_paragraph()
    title.style = "Title"
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    title.add_run("Video Agent Studio\n完整工作流说明文档")

    subtitle = document.add_paragraph()
    subtitle.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = subtitle.add_run(
        "面向项目理解：从用户点击按钮到产出 artifact 的完整时序、Agent 提示词装配、读写边界与存储逻辑"
    )
    run.italic = True
    run.font.size = Pt(10.5)
    run.font.color.rgb = RGBColor(90, 90, 90)


def add_bullets(document: Document, items: list[str]) -> None:
    for item in items:
        paragraph = document.add_paragraph(style="List Bullet")
        paragraph.add_run(item)


def add_numbered(document: Document, items: list[str]) -> None:
    for item in items:
        paragraph = document.add_paragraph(style="List Number")
        paragraph.add_run(item)


def add_code_block(document: Document, text: str) -> None:
    paragraph = document.add_paragraph(style="No Spacing")
    run = paragraph.add_run(text.rstrip())
    run.font.name = "Courier New"
    run.font.size = Pt(8.7)


def add_agent_table(document: Document) -> None:
    document.add_heading("一、9 个 Agent 的职责、输入、输出与读写边界", level=1)
    document.add_paragraph(
        "下面这张表不是抽象想象，而是基于当前仓库真实实现整理出来的“运行语义表”。其中 Agent5 现在主要以规则审核器存在，"
        "Agent8/9 负责生成媒体 artifact，Agent1 主要承担人工对齐与失败升级后的重规划。"
    )

    headers = ["Agent", "主职责", "主要读取", "主要写入", "存储落点", "关键约束"]
    rows = [
        (
            "Agent1",
            "对齐需求、产出配置建议；当连续失败时负责重规划。",
            "项目灵感、当前 config、失败总结、review 历史摘要。",
            "配置建议、重规划建议、stage snapshot。",
            "task.outputJson/outputMarkdown；projectStageSnapshots(agent1_output)；config 由人工确认页单独写入。",
            "是人工门，不直接推进最终 artifact；失败升级后才自动介入。",
        ),
        (
            "Agent2",
            "调研：事实、洞察、建议。",
            "source idea、active config、上游上下文。",
            "research markdown/json、task 记录、snapshot。",
            "agentTasks；projectStageSnapshots(research)。",
            "必须满足调研结构和事实数量阈值。",
        ),
        (
            "Agent3",
            "把调研整理成故事线与结构卡片。",
            "source idea、config、research snapshot。",
            "storyline markdown/json、task 记录、snapshot。",
            "agentTasks；projectStageSnapshots(storyline)。",
            "必须输出足够数量的 storyline cards。",
        ),
        (
            "Agent4",
            "形成可进入视觉流程的最终脚本。",
            "config、research/storyline snapshot。",
            "script markdown/json、task 记录、snapshot。",
            "agentTasks；projectStageSnapshots(script)。",
            "脚本要足够细，供 Agent6 拆 scene/shot。",
        ),
        (
            "Agent5",
            "审核与质量门控。",
            "task.outputMarkdown、task.outputJson。",
            "review records、revision brief。",
            "reviewRecords。",
            "当前实现主要是规则引擎，不是第二个自由生成 Agent。",
        ),
        (
            "Agent6",
            "把脚本拆成 scene 和 shot 结构。",
            "script snapshot、scriptSections、scriptExcerpt。",
            "scenes/shots 结构、task 记录、snapshot。",
            "agentTasks；projectStageSnapshots(shots)；sceneRepository；shotRepository。",
            "不能随意压缩场景；必须保留脚本信息密度。",
        ),
        (
            "Agent7",
            "逐 shot 扩写视觉描述。",
            "script snapshot、当前 shot chunk、scene outline、Agent6 结果。",
            "visualPrompt、lighting、cameraMotion、compositionNotes 等。",
            "agentTasks；shotRepository；projectStageSnapshots(storyboard_spec)。",
            "逐 shot 处理；输出数量必须与 shotCount 精确匹配。",
        ),
        (
            "Agent8",
            "生成 storyboard image。",
            "shotRepository 中的结构化 shot 字段、visualPrompt。",
            "artifact group/version、图片资源、active version 切换。",
            "artifactGroups/artifactVersions/shotAssetBindings；.data/generated-assets。",
            "不回写正文 artifact；遇锁定版本不能覆盖。",
        ),
        (
            "Agent9",
            "生成 video clip。",
            "project title、shot title 或局部 regenerate prompt。",
            "video artifact group/version、媒体资源、active version 切换。",
            "artifactGroups/artifactVersions/shotAssetBindings；.data/generated-assets 或远端 URL。",
            "同样不改写文本 artifact；遇锁定版本不能覆盖。",
        ),
    ]

    table = document.add_table(rows=1, cols=len(headers))
    table.style = "Table Grid"
    header_cells = table.rows[0].cells
    for index, header in enumerate(headers):
        header_cells[index].text = header
        set_cell_shading(header_cells[index], "D9EAF7")

    for row in rows:
        cells = table.add_row().cells
        for index, value in enumerate(row):
            cells[index].text = value


def add_prompt_logic_section(document: Document) -> None:
    document.add_heading("二、Agent 内部提示词规划方式", level=1)
    document.add_paragraph(
        "这里讲的是项目内 Agent 运行时的 prompt 装配逻辑，而不是外部助手平台的隐藏系统提示词。项目内真实逻辑位于 "
        "packages/agents/src/runtime/execute-agent.ts。"
    )
    add_bullets(
        document,
        [
            "每次执行都会进入 executeAgent(input)。如果当前 model binding 是 mock，则直接走 executeAgentStub，返回预制结构化内容。",
            "如果是非 mock provider，则系统会拼接 instructions 和 prompt 两层内容，再要求模型按统一 JSON schema 返回。",
            "统一输出 contract 是 outputSummary、outputMarkdown、outputJson 三字段；这样后续 review、snapshot、下游 workflow 都可以稳定消费。",
            "instructions 层由四部分组成：基础身份说明、Agent-specific requirements、项目设置页可编辑的 roleDefinition/systemPrompt/developerPrompt、reviewRules。",
            "prompt 层会携带 project title、source idea、Current active config、Current task payload、Revision brief，以及 Workflow context JSON。",
            "当模型没有返回合格 JSON 时，系统会依次进入 retry mode 和 formatting recovery mode，尽量把自由文本拉回结构化结果。",
        ],
    )

    document.add_paragraph("当前仓库里，Agent2/3/4/6/7 都有明确的结构化输出要求。典型要求如下：")
    add_bullets(
        document,
        [
            "Agent2 必须输出“调研目标 / 事实依据 / 洞察与风险 / 创作建议”，并在 outputJson 中带 facts、insights、recommendations、openQuestions。",
            "Agent3 必须输出“故事总纲 / 结构卡片 / 节奏与情绪设计”，并提供至少 4 张 storyline cards。",
            "Agent4 必须输出“片子定位 / 成片目标 / 分段脚本”，并提供至少 3 个 scriptSections。",
            "Agent6 必须输出“场景拆解 / Shot 列表 / 连贯性说明”，scene/shot 数量由脚本内容决定，而不是写死。",
            "Agent7 必须输出“扩写原则 / Shot 扩写结果”，且 shotExpansions 数量必须与输入 shotCount 完全一致。",
        ],
    )

    document.add_paragraph("提示词装配顺序示意：")
    add_code_block(
        document,
        """
executeAgent
  -> buildInstructions
     -> agent objective
     -> stage name
     -> agent-specific requirements
     -> roleDefinition / systemPrompt / developerPrompt / reviewRules
  -> buildPrompt
     -> projectTitle / sourceIdea
     -> Current active config
     -> Current task payload
     -> revisionBrief
     -> Workflow context JSON
  -> provider.generate(json_schema)
  -> parse / retry / formatting recovery
        """,
    )

    document.add_paragraph(
        "很重要的一点：设置页虽然还允许编辑 inputSchema、outputSchema、allowedTools、contextScope、writeScope、deleteScope，"
        "但按照当前代码，这些字段还没有形成强制执行层；真正进入运行时的是 roleDefinition、systemPrompt、developerPrompt、reviewRules。"
    )


def add_sequence_section(document: Document) -> None:
    document.add_heading("三、从用户给出灵感到产出 artifact 的完整时序图", level=1)
    document.add_paragraph(
        "这张时序图把“用户点击按钮”后的实际系统路径展开到了 API、workflow、review、DB、artifact 和前端工作台层。"
    )
    add_code_block(
        document,
        r"""
[用户输入灵感]
    |
    v
Project Create
    -> projectRepository.create
    -> .data/video-agent-studio.json 新建 project
    |
    v
进入 Agent1 对齐页
    -> 读取当前 active config / overview
    -> 用户编辑配置
    -> POST /api/projects/{projectId}/config/versions
    -> configRepository.createDraft
    -> POST /api/projects/{projectId}/config/versions/{configVersionId}/confirm
    -> configRepository.confirm
    |
    v
点击“确认配置并进入第一部分”
    -> POST /api/projects/{projectId}/workflows/script/run
    -> runScriptWorkflow
       -> runRepository.create(run: script)
       -> settingsRepository.ensureDefaults
       -> 检查 Agent2/3/4 provider binding
       -> 顺序执行 Agent2 -> Agent3 -> Agent4
          -> runReviewedNode
             -> createTaskRecord
                -> executeAgent / executeAgentStub
                -> taskRepository.create
             -> reviewTaskOutput (Agent5规则审核)
             -> pass: task approved + createStageSnapshot
             -> revise: 自动重试，最多 3 轮
             -> 连续 3 次失败:
                -> createFailureSummaryDoc
                -> Agent1 重规划
                -> 第 4 轮 recovery
       -> 全部通过后 run.currentNode = script_user_confirm
       -> 前端显示“确认脚本进入第二部分”
    |
    v
点击“确认脚本进入第二部分”
    -> POST /api/projects/{projectId}/script/confirm
    -> project.currentStage = storyboard
    |
    v
点击“运行第二部分 / storyboard”
    -> POST /api/projects/{projectId}/workflows/storyboard/run
    -> runStoryboardWorkflow
       -> runRepository.create(run: storyboard)
       -> Agent6:
          -> 读取 script snapshot
          -> 生成 scenes/shots 结构
          -> syncStoryboardStructure
          -> sceneRepository.replaceForProject
          -> shotRepository.create/update/delete
       -> Agent7:
          -> 按 chunk(当前实现是 1 shot 一批) 逐 shot 扩写
          -> 把 visualPrompt / lighting / cameraMotion / compositionNotes / styleNotes 写回 shotRepository
       -> Agent8:
          -> 遍历所有 shot
          -> buildStoryboardImagePrompt
          -> imageProvider.generate
          -> materializeMediaOutput
          -> artifactRepository.createGroup/createVersion
          -> artifactRepository.activateVersion
          -> shotAssetBindings 绑定 storyboard_main
    |
    v
工作台出现 storyboard artifact
    -> Project overview / Shot workbench / Timeline board 都能读到当前 active version
    |
    v
点击“运行第三部分 / video”
    -> POST /api/projects/{projectId}/workflows/video/run
    -> runVideoWorkflow
       -> runRepository.create(run: video)
       -> 遍历 shot
       -> createVideoProvider.generate
       -> materializeMediaOutput
       -> artifactRepository.createVersion
       -> artifactRepository.activateVersion
       -> shotAssetBindings 绑定 video_main
    |
    v
最终产出
    -> shot 对应 storyboard image artifact
    -> shot 对应 video clip artifact
    -> overview / timeline / logs 中持续可见
        """,
    )

    document.add_paragraph("如果用户不是全流程运行，而是在时间线中做局部重生成，时序会变成：")
    add_numbered(
        document,
        [
            "点击某个 shot 的“重生成分镜图”或“重生成视频”。",
            "前端先调用 impact-analysis，分析相邻 shot、关联 storyboard/video/key element、锁定状态和影响级别。",
            "如果影响中高或存在锁冲突，需要用户确认说明；确认后再真正发起 regenerate API。",
            "API 先创建一个 status=pending 的新版本，再用 runDetached 在后台异步生成。",
            "后台成功则把 pending 版本更新为 generated，并切换 active version；失败则把该版本标成 failed。",
        ],
    )


def add_read_write_storage_section(document: Document) -> None:
    document.add_heading("四、每一步的阅读、撰写、存储逻辑", level=1)
    document.add_paragraph(
        "这一部分回答三个问题：谁在读什么、谁在写什么、写到哪里。"
    )

    document.add_heading("4.1 读取逻辑", level=2)
    add_bullets(
        document,
        [
            "所有文本 Agent 在 createTaskRecord 时都会读取 active config、project overview、当前 agent profile、当前 model binding。",
            "overview 中最重要的是 activeSnapshots；它决定当前 Agent 能继承哪些上游结果。",
            "脚本链通常基于 config 和上游 task snapshot 运行；storyboard 链会显式读取 script snapshot；video 链更多直接读取 shot 和 artifact 状态。",
            "Agent6 读取 scriptSections 和 scriptExcerpt；Agent7 读取 shot chunk、scene outline 和 script snapshot；Agent8 读取 shotRepository 里已经结构化的字段，而不是重新读整段脚本。",
            "评审器 Agent5 不读外部上下文，主要读 task.outputMarkdown 与 task.outputJson。",
        ],
    )

    document.add_heading("4.2 撰写逻辑", level=2)
    add_bullets(
        document,
        [
            "文本 Agent 先写 task 记录，再在审核通过后写 stage snapshot；也就是说 task 是完整执行历史，snapshot 是当前激活态摘要。",
            "Agent6 不只是写 task，它还会把结构真正同步到 sceneRepository 和 shotRepository，所以它是文本层和结构层之间的桥。",
            "Agent7 的核心价值是把 shot 从“可读描述”升级成“可生成描述”，因此它重点回写 visualPrompt、lighting、cameraMotion、compositionNotes、styleNotes。",
            "Agent8/9 不改写脚本正文；它们写的是 artifact group、artifact version，以及本地生成媒体文件或远端 URL。",
            "局部 regenerate 不会覆盖旧版本，而是永远追加一个新版本，再切 activeVersionId。",
        ],
    )

    document.add_heading("4.3 存储逻辑", level=2)
    add_bullets(
        document,
        [
            "项目主数据：.data/video-agent-studio.json。",
            "生成图片/视频的本地媒体文件：.data/generated-assets/。",
            "配置版本：projectConfigVersions。",
            "执行历史：workflowRuns、agentTasks、reviewRecords、taskEvents。",
            "当前有效摘要：projectStageSnapshots。",
            "镜头结构：scenes、shots。",
            "多媒体版本：artifactGroups、artifactVersions、shotAssetBindings。",
            "锁与冲突确认：objectLocks、userAnnotations(conflict_resolution)。",
        ],
    )

    document.add_paragraph("推荐你把数据层理解为三层：")
    add_numbered(
        document,
        [
            "Execution Layer：run / task / review / event，回答“系统当时做了什么”。",
            "Working State Layer：config / snapshot / scene / shot，回答“当前工作台现在应该长什么样”。",
            "Artifact Layer：group / version / binding / lock，回答“最终素材有哪些版本，当前生效的是哪一个”。",
        ],
    )


def add_agent_by_agent_detail(document: Document) -> None:
    document.add_heading("五、从灵感出发，每个 Agent 是如何开展工作的", level=1)
    sections = [
        (
            "Agent1",
            [
                "主要承担项目理解与配置对齐，而不是直接生成最终脚本。",
                "在正常主流程中，Agent1 更多出现在人工对齐页面；在失败升级场景中，它会根据 failure summary 给出重规划建议。",
                "因此 Agent1 更像 workflow governor，而不是常规内容生产节点。",
            ],
        ),
        (
            "Agent2 -> Agent3 -> Agent4",
            [
                "这是把“灵感”变成“可拍脚本”的文本链路。",
                "Agent2 把命题拆成事实、洞察、建议；Agent3 把调研重组成叙事结构；Agent4 把结构落成分段脚本。",
                "每经过一个节点，系统都会先记录 task，再执行 review；通过的内容才会晋升为 active snapshot。",
            ],
        ),
        (
            "Agent6 -> Agent7",
            [
                "这是把“脚本语言”翻译成“视觉生产语言”的关键桥梁。",
                "Agent6 决定场景数、shot 数、shot 顺序，并直接塑造 Part2/Part3 工作台的数据骨架。",
                "Agent7 进一步把每个 shot 扩成 image generation 可消费的结构，特别强调主体、动作、文字排版、构图、阅读顺序。",
            ],
        ),
        (
            "Agent8 -> Agent9",
            [
                "这是 artifact 生产层。",
                "Agent8 根据结构化 shot prompt 生成 storyboard image；Agent9 生成 video clip。",
                "这两个 Agent 的成功输出不是 markdown，而是 artifact version 与其对应的媒体内容。",
            ],
        ),
    ]

    for heading, bullets in sections:
        document.add_heading(heading, level=2)
        add_bullets(document, bullets)


def add_control_plane_section(document: Document) -> None:
    document.add_heading("六、控制平面：锁、影响分析、冲突确认、版本激活", level=1)
    add_bullets(
        document,
        [
            "这个产品不是点一下就盲目重算，它会先做 impact analysis，评估受影响对象和风险级别。",
            "如果对象本身被锁、相邻 shot 会被影响、当前 active storyboard/video/key element 已锁定，就会升级成 requiresUserConfirmation。",
            "锁定的对象可读但不能自动覆盖；这与项目 AGENTS.md 约束完全一致。",
            "版本切换不是覆盖文件，而是切换 artifactGroup.activeVersionId。",
            "因此，工作台默认展示的是当前 active version，但历史版本会一直保留，满足 traceability。",
        ],
    )

    document.add_paragraph("控制平面上的一个典型动作链路：")
    add_code_block(
        document,
        """
用户点击“重生成分镜图”
  -> ControlPlaneActionButton
  -> /api/projects/{projectId}/impact-analysis
  -> controlPlaneRepository.analyzeImpact
  -> 如果需要确认:
     -> 创建 conflict_resolution annotation
     -> 用户填写 resolution
     -> /api/conflicts/{conflictId}/resolve
  -> 真正发起 regenerate API
  -> 创建新的 pending artifact version
  -> 后台生成
  -> 成功后 activateVersion
        """,
    )


def add_runtime_reality_section(document: Document) -> None:
    document.add_heading("七、你需要特别知道的当前 MVP 现实", level=1)
    add_bullets(
        document,
        [
            "当前仓库最成熟的是 workflow 骨架、版本机制、锁机制、局部重生成、结构化 prompt 装配。",
            "当前 worker 进程不是独立调度中心；真正的 workflow 调用主要还是由 web API 直接触发。",
            "如果 provider 还是 mock，大部分文本 Agent 会直接走 stub 输出，因此你看到的是“可审计的占位流程”，不是最终模型能力。",
            "Agent5 现在主要是 deterministic review service，不是自由评论员式大模型审核。",
            "视频 prompt 侧目前比 storyboard prompt 侧更轻，说明系统当前重心还是先把结构化工作台与分镜层打稳。",
        ],
    )


def add_source_index(document: Document) -> None:
    document.add_heading("八、关键源码入口索引", level=1)
    source_items = [
        "prompt 装配：packages/agents/src/runtime/execute-agent.ts",
        "统一执行-审核-重试内核：apps/worker/src/workflows/reviewed-node.ts",
        "脚本工作流：apps/worker/src/workflows/run-script.workflow.ts",
        "分镜工作流：apps/worker/src/workflows/run-storyboard.workflow.ts",
        "视频工作流：apps/worker/src/workflows/run-video.workflow.ts",
        "状态机：packages/workflow-engine/src/core/state-machine.ts",
        "审核规则：packages/review-service/src/index.ts",
        "分镜 prompt builder：packages/shared/src/lib/storyboard-prompt.ts",
        "artifact 版本与媒体落地：packages/artifact-service/src/index.ts",
        "设置中心与 binding 默认值：packages/db/src/repositories/settings.repository.ts",
        "控制平面影响分析：packages/db/src/repositories/control-plane.repository.ts",
        "前端控制面壳：apps/web/components/project-shell.tsx",
        "横向时间线工作台：apps/web/components/timeline-board.tsx",
        "Shot 编辑工作台：apps/web/components/shot-workbench.tsx",
        "局部分镜重生成：apps/web/app/api/shots/[shotId]/storyboard/regenerate/route.ts",
        "局部视频重生成：apps/web/app/api/shots/[shotId]/video/regenerate/route.ts",
    ]
    add_bullets(document, source_items)


def add_summary(document: Document) -> None:
    document.add_heading("九、最终理解模型", level=1)
    document.add_paragraph(
        "如果只用一句话概括这个系统：Video Agent Studio 当前是一个“以工作流控制面为核心的视频生产 MVP”，"
        "不是一个自由对话机器人。它把用户灵感依次转成配置、调研、故事线、脚本、scene/shot、storyboard artifact 和 video artifact；"
        "中间每一步都有 task 记录、审核门控、快照激活、版本追踪、锁与冲突控制。"
    )
    document.add_paragraph(
        "建议你以后理解它时，不要把它看成“9 个 agent 在群聊里互相讲话”，而要把它看成："
        "一个 workflow engine 驱动的、带结构化 prompt contract 的、可审计的生产流水线。"
    )


def build_document() -> Path:
    OUTPUT_DIR.mkdir(parents=True, exist_ok=True)
    document = Document()
    set_document_defaults(document)
    section = document.sections[0]
    section.top_margin = Inches(0.75)
    section.bottom_margin = Inches(0.75)
    section.left_margin = Inches(0.7)
    section.right_margin = Inches(0.7)

    add_title(document)
    document.add_paragraph("")
    add_agent_table(document)
    add_prompt_logic_section(document)
    add_sequence_section(document)
    add_read_write_storage_section(document)
    add_agent_by_agent_detail(document)
    add_control_plane_section(document)
    add_runtime_reality_section(document)
    add_source_index(document)
    add_summary(document)
    document.save(OUTPUT_PATH)
    return OUTPUT_PATH


if __name__ == "__main__":
    path = build_document()
    print(path)
